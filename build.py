#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build.py  -  Real Estate Analytics  (Departamentos Comparativa)

Lee el Excel que esté en ./data/, lo limpia, geocodifica las direcciones
(con caché en disco para no re-pedir lo ya resuelto), calcula métricas
derivadas + sub-scores normalizados, y emite ./data.js que consume el sitio.

Diseño:
- ADAPTATIVO: las columnas se detectan por nombre (sin acentos / sin importar
  mayúsculas) y la fila de encabezados se detecta automáticamente (tolera filas
  de título arriba). Las columnas no reconocidas se arrastran a "extra" y se
  muestran en el detalle del sitio.
- IDEMPOTENTE: sólo geocodifica direcciones que no estén en el caché.
- DETERMINISTA: si el caché está completo, no necesita internet.

Uso:
    python build.py                  # build normal
    python build.py --no-geocode     # nunca llama a la red (sólo caché)
    python build.py --rate 1150      # fija el tipo de cambio USD/ARS por defecto
"""
from __future__ import annotations
import argparse
import copy
import datetime
import json
import math
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalá con:  pip install -r requirements.txt")

# --------------------------------------------------------------------------- #
# Configuración
# --------------------------------------------------------------------------- #
ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
CACHE_FILE = DATA_DIR / "geocode_cache.json"
SUBTE_FILE = DATA_DIR / "subte_stations.json"
SUBTE_LINES_FILE = DATA_DIR / "subte_lines.json"
OUT_FILE = ROOT / "data.js"              # PÚBLICO: se commitea y se publica (sin precios de negociación)
PRIVATE_FILE = ROOT / "data.private.js"  # LOCAL: datos completos; gitignoreado, sólo para preview local

# Campos sensibles (estrategia de negociación) que NUNCA se publican.
SENSITIVE_FIELDS = ["precio_max", "negociar_a", "usd_m2_max", "dif_vs_tope", "margen_neg_pct"]

USD_ARS_RATE_DEFAULT = 1200          # tipo de cambio editable desde la UI
CABA_VIEWBOX = "-58.531,-34.526,-58.335,-34.705"   # left,top,right,bottom
NOMINATIM = "https://nominatim.openstreetmap.org/search"
USER_AGENT = "real-estate-analytics/1.0 (github-pages build script)"

# Mapeo canónico -> posibles encabezados (normalizados: sin acento, minúscula).
# El match es por igualdad exacta del header normalizado; si varios coinciden,
# gana el alias más largo. Así "usd/m2 precio max" no pisa a "usd/m2".
FIELD_ALIASES = {
    "top":          ["top", "ranking", "orden", "#", "puesto"],
    "motivo":       ["motivo", "notas", "observaciones", "comentario", "comentarios", "descripcion"],
    "ambientes":    ["amb", "amb.", "ambientes", "ambs"],
    "link_aviso":   ["link aviso", "aviso", "link", "url", "publicacion"],
    "direccion":    ["direccion", "domicilio", "ubicacion", "calle"],
    "maps":         ["google maps", "maps", "mapa"],
    "precio_max":   ["precio maximo aceptable", "precio max aceptable", "precio max", "maximo aceptable", "tope"],
    "negociar_a":   ["negociar a", "negociar", "oferta"],
    "precio":       ["precio (usd)", "precio usd", "precio", "valor", "precio de lista"],
    "p25":          ["25% del precio", "25%", "anticipo 25%"],
    "p75":          ["75% del precio", "75%", "credito 75%"],
    "expensas":     ["expensas (ars)", "expensas ars", "expensas"],
    "m2":           ["m2 totales", "m totales", "m2", "metros", "superficie", "sup", "m2 total"],
    "usd_m2":       ["usd/m2", "usd m2", "usd por m2", "precio m2"],
    "usd_m2_max":   ["usd/m2 precio max", "usd m2 precio max"],
    "antiguedad":   ["antiguedad (anos)", "antiguedad", "antig", "anos", "edad"],
    "piso":         ["piso", "nivel"],
    "orientacion":  ["orientacion"],
    "posicion":     ["posicion", "ubicacion en el piso"],
    "balcon":       ["balcon/terraza", "balcon", "terraza", "exterior", "balcon / terraza"],
    "apto_credito": ["apto credito", "credito", "apto cred"],
    "lat":          ["lat", "latitud", "latitude"],
    "lng":          ["lng", "lon", "long", "longitud", "longitude"],
}

# Columnas cuya ausencia degrada métricas (no aborta, pero avisa).
KEY_FIELDS = ["precio", "m2", "top"]

NUMERIC_FIELDS = {"top", "ambientes", "precio_max", "negociar_a", "precio", "p25",
                  "p75", "expensas", "m2", "usd_m2", "usd_m2_max", "antiguedad",
                  "piso", "lat", "lng"}

MISSING_TOKENS = {"", "?", "-", "n/a", "na", "s/d", "sd", "nan", "none"}

ALL_ALIASES = {a for als in FIELD_ALIASES.values() for a in als}


# --------------------------------------------------------------------------- #
# Helpers de normalización / parsing
# --------------------------------------------------------------------------- #
def strip_accents(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")


def norm(s) -> str:
    """minúscula, sin acentos, espacios colapsados."""
    if s is None:
        return ""
    return re.sub(r"\s+", " ", strip_accents(str(s)).lower()).strip()


def parse_num(v):
    """Devuelve float o None. Tolera '?', '$', 'USD' y separadores de miles AR/US.
    Ej: 'USD 95.000'->95000, '$ 90,000'->90000, '1.200.000'->1200000,
        '75.000,50'->75000.5, '1583.33'->1583.33."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    s = str(v).strip()
    if norm(s) in MISSING_TOKENS:
        return None
    neg = s.lstrip().startswith("-")
    s = re.sub(r"(?i)(usd|ars|u\$s|\$|m2|m²)", "", s)
    s = re.sub(r"[^\d,.]", "", s)
    if s in ("", ".", ","):
        return None
    has_dot, has_comma = "." in s, "," in s
    if has_dot and has_comma:                       # el último separador es el decimal
        if s.rfind(",") > s.rfind("."):             # 75.000,50  -> coma decimal
            s = s.replace(".", "").replace(",", ".")
        else:                                       # 75,000.50  -> punto decimal
            s = s.replace(",", "")
    elif has_comma:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", s):    # 90,000 -> miles
            s = s.replace(",", "")
        else:                                       # 1,5 -> decimal
            s = s.replace(",", ".")
    elif has_dot:
        if s.count(".") > 1 or re.fullmatch(r"\d{1,3}(\.\d{3})+", s):  # 1.200.000 / 95.000 -> miles
            s = s.replace(".", "")
        # else: punto decimal normal (1583.33) -> se deja
    try:
        val = float(s)
        return -val if neg else val
    except ValueError:
        print(f"  ! No pude parsear número: {v!r}")
        return None


def parse_str(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return None if norm(s) in MISSING_TOKENS else s


def safe_url(v):
    """Devuelve la URL sólo si es http(s); si no, None (evita javascript:/data:)."""
    s = parse_str(v)
    return s if (s and re.match(r"(?i)^https?://", s)) else None


# --------------------------------------------------------------------------- #
# Geocodificación (Nominatim) con caché
# --------------------------------------------------------------------------- #
def load_json(path, default):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception as e:
            print(f"  ! No pude leer {path.name}: {e}")
    return default


# Abreviaturas comunes de calles AR -> nombre completo (Nominatim matchea mejor el completo).
STREET_ABBR = {
    "av": "Avenida", "avda": "Avenida", "gral": "General", "grl": "General",
    "cnel": "Coronel", "tte": "Teniente", "pres": "Presidente", "pte": "Presidente",
    "dr": "Doctor", "dra": "Doctora", "gob": "Gobernador", "alte": "Almirante",
    "brig": "Brigadier", "ing": "Ingeniero", "arq": "Arquitecto", "cap": "Capitan",
    "sgto": "Sargento", "prof": "Profesor", "mons": "Monsenor", "pje": "Pasaje",
}


def expand_abbr(text: str) -> str:
    """Expande tokens abreviados con punto: 'Gral. Artigas' -> 'General Artigas'."""
    out = []
    for tok in text.split():
        if tok.endswith("."):
            k = norm(tok).rstrip(".")
            if k in STREET_ABBR:
                out.append(STREET_ABBR[k]); continue
        out.append(tok)
    return " ".join(out)


def derive_query(direccion: str):
    """De 'Flores Artigas 46, Flores, Capital Federal' -> ('Artigas 46', 'Flores')."""
    s = direccion.strip()
    s = re.sub(r"(?i),?\s*(capital federal|caba|c\.?a\.?b\.?a\.?|"
               r"ciudad autonoma de buenos aires|buenos aires)\s*$", "", s).strip().strip(",")
    parts = [p.strip() for p in s.split(",") if p.strip()]
    street = parts[0] if parts else s
    barrio = parts[1] if len(parts) > 1 else None
    if barrio and norm(street).startswith(norm(barrio) + " "):  # quita barrio duplicado en la calle
        street = street[len(barrio):].strip()
    return expand_abbr(street), barrio


def geocode(direccion: str, cache: dict, allow_net: bool):
    key = norm(direccion)
    if key in cache and cache[key].get("lat") is not None:
        hit = dict(cache[key]); hit["from_cache"] = True
        return hit
    if not allow_net:
        return {"lat": None, "lng": None, "barrio": None, "source": "miss", "from_cache": False}

    street, barrio = derive_query(direccion)
    q = ", ".join(filter(None, [street, barrio, "Ciudad Autónoma de Buenos Aires", "Argentina"]))
    params = {"format": "json", "limit": "1", "countrycodes": "ar",
              "viewbox": CABA_VIEWBOX, "bounded": "1", "addressdetails": "1", "q": q}
    url = NOMINATIM + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    print(f"  geocode -> {q}")
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=25) as r:
                res = json.load(r)
            time.sleep(1.1)            # política de uso Nominatim: <= 1 req/s
            if res:
                a = res[0].get("address", {})
                rec = {
                    "lat": float(res[0]["lat"]), "lng": float(res[0]["lon"]),
                    "barrio": a.get("suburb") or a.get("neighbourhood")
                              or a.get("city_district") or barrio,
                    "source": "nominatim",
                    "display": res[0].get("display_name", "")[:120],
                }
                cache[key] = rec
                out = dict(rec); out["from_cache"] = False
                return out
            break
        except Exception as e:
            print(f"    intento {attempt+1} falló: {e}")
            time.sleep(2)
    rec = {"lat": None, "lng": None, "barrio": barrio, "source": "failed"}
    cache[key] = rec
    out = dict(rec); out["from_cache"] = False
    return out


def haversine_m(lat1, lon1, lat2, lon2):
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(min(1, math.sqrt(a)))


def nearest_subte(lat, lng, stations):
    if lat is None or lng is None or not stations:
        return None, None
    best, bestd = None, float("inf")
    for st in stations:
        d = haversine_m(lat, lng, st["lat"], st["lon"])
        if d < bestd:
            bestd, best = d, st
    return best, bestd


# --------------------------------------------------------------------------- #
# Detección de fila de encabezados + columnas
# --------------------------------------------------------------------------- #
def find_header_row(rows):
    """Devuelve (idx_0based, [(col, val), ...]) de la 1ª fila con >=2 aliases conocidos.
    Tolera filas de título mergeadas arriba. Fallback: 1ª fila no vacía."""
    best = None  # (hits, idx, cells)
    for i, row in enumerate(rows[:20]):
        cells = [(c.column, c.value) for c in row]
        hits = sum(1 for _, v in cells if v is not None and norm(v) in ALL_ALIASES)
        if hits >= 2 and (best is None or hits > best[0]):
            best = (hits, i, cells)
    if best:
        return best[1], best[2]
    for i, row in enumerate(rows):       # fallback
        cells = [(c.column, c.value) for c in row]
        if any(v is not None and str(v).strip() for _, v in cells):
            return i, cells
    return 0, [(c.column, c.value) for c in rows[0]]


def resolve_columns(headers):
    """headers: lista de (idx, texto). Devuelve {field: idx} y lista de extras."""
    norm_headers = [(i, norm(h)) for i, h in headers if h is not None and str(h).strip()]
    mapping, used_idx = {}, set()
    for field, aliases in FIELD_ALIASES.items():
        best = None  # (alias_len, idx)
        for i, h in norm_headers:
            if i in used_idx:
                continue
            for al in aliases:
                if h == al and (best is None or len(al) > best[0]):
                    best = (len(al), i)
        if best:
            mapping[field] = best[1]
            used_idx.add(best[1])
    extras = [(i, dict(headers)[i]) for i, _ in norm_headers if i not in used_idx]
    return mapping, extras


# --------------------------------------------------------------------------- #
# Scoring helpers
# --------------------------------------------------------------------------- #
def minmax_norm(values, invert=False):
    """Lista de (None|float) -> lista 0..1 (None imputado a 0.5 = neutro)."""
    nums = [v for v in values if v is not None]
    if not nums:
        return [0.5] * len(values)
    lo, hi = min(nums), max(nums)
    span = hi - lo
    out = []
    for v in values:
        if v is None or span == 0:
            out.append(0.5)
        else:
            n = (v - lo) / span
            out.append(1 - n if invert else n)
    return out


ORIENT_SCORE = {"n": 1.0, "ne": 0.88, "no": 0.88, "e": 0.72, "o": 0.72,
                "se": 0.55, "so": 0.55, "s": 0.38}
OUTDOOR_SCORE = {"terraza": 1.0, "balcon": 0.8, "patio": 0.55, "no": 0.2}
POSICION_SCORE = {"frente": 1.0, "contra frente": 0.7, "contrafrente": 0.7,
                  "interno": 0.5, "lateral": 0.6}

RE_REFACCION = re.compile(r"(?i)\b(arreglar|refacc|reciclar|validar|"
                          r"arregl|renovar|pintar|cambiar)")
RE_TRANSPORTE = re.compile(r"(?i)(subte|colectivo|tren|metrobus|estacion|plaza|parque|"
                           r"cuadra)")


# --------------------------------------------------------------------------- #
# Build
# --------------------------------------------------------------------------- #
def find_excel():
    cands = sorted(DATA_DIR.glob("*.xlsx")) + sorted(DATA_DIR.glob("*.xlsm"))
    cands = [c for c in cands if not c.name.startswith("~$")]
    if not cands:
        sys.exit(f"No encontré ningún .xlsx en {DATA_DIR}")
    pref = [c for c in cands if "comparativa" in norm(c.stem)]
    return (pref or cands)[0]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-geocode", action="store_true", help="No llamar a la red")
    ap.add_argument("--rate", type=float, default=USD_ARS_RATE_DEFAULT)
    ap.add_argument("--full", action="store_true",
                    help="Además emite data.private.js con los precios de negociación (sólo para preview local)")
    args = ap.parse_args()

    excel = find_excel()
    print(f"Excel:  {excel.name}")
    wb_v = openpyxl.load_workbook(excel, data_only=True)
    wb_f = openpyxl.load_workbook(excel, data_only=False)  # para hipervínculos/fórmulas
    ws_v = wb_v[wb_v.sheetnames[0]]
    ws_f = wb_f[wb_f.sheetnames[0]]

    rows = list(ws_v.iter_rows())
    if not rows:
        sys.exit("La hoja está vacía: no hay encabezados ni datos.")

    header_idx, headers = find_header_row(rows)
    if header_idx > 0:
        print(f"Encabezados detectados en la fila {header_idx + 1} (salté {header_idx} fila/s de título).")
    mapping, extras = resolve_columns(headers)
    print(f"Columnas detectadas: {', '.join(sorted(mapping))}")
    if extras:
        print(f"Columnas extra (se muestran en el detalle): {[e[1] for e in extras]}")
    if "direccion" not in mapping:
        sys.exit("No encontré la columna de dirección; no puedo geocodificar.")
    for k in KEY_FIELDS:
        if k not in mapping:
            print(f"  ! Aviso: no detecté columna para '{k}'; algunas métricas quedarán en s/d.")

    cache = load_json(CACHE_FILE, {})
    stations = load_json(SUBTE_FILE, [])
    subte_lines = load_json(SUBTE_LINES_FILE, [])
    # saneo de datos OSM (colour válido CSS, ref/name como texto) — defensa anti-inyección
    for L in subte_lines:
        c = str(L.get("colour") or "")
        L["colour"] = c if re.match(r"^#[0-9A-Fa-f]{3,8}$|^[a-z]+$", c) else "#888"
        L["ref"] = parse_str(L.get("ref")) or "?"
        L["name"] = parse_str(L.get("name")) or ""
    print(f"Estaciones de subte cargadas: {len(stations)}  ·  líneas: {len(subte_lines)}")

    def cell_text(rowcells, field):
        idx = mapping.get(field)
        return rowcells[idx - 1].value if idx else None

    props = []
    for j in range(header_idx + 1, len(rows)):
        rowcells = rows[j]
        excel_row = j + 1
        # saltar sólo filas COMPLETAMENTE vacías
        if not any(c.value is not None and str(c.value).strip() for c in rowcells):
            continue

        direccion = parse_str(cell_text(rowcells, "direccion"))
        rec = {"row": excel_row}
        for field in mapping:
            raw = cell_text(rowcells, field)
            rec[field] = parse_num(raw) if field in NUMERIC_FIELDS else parse_str(raw)

        # hipervínculo real del aviso (validado a http/https)
        li = mapping.get("link_aviso")
        hl = ws_f.cell(row=excel_row, column=li).hyperlink if li else None
        rec["link_aviso_url"] = safe_url(hl.target) if hl else None

        # columnas extra (serialización segura)
        rec["extra"] = {}
        for i, name in extras:
            val = rowcells[i - 1].value
            v = val if isinstance(val, (int, float, bool)) else parse_str(val)
            if v is not None and v != "":
                rec["extra"][str(name)] = v

        # geocode (lat/lng del Excel tienen prioridad si existen)
        lat, lng = rec.get("lat"), rec.get("lng")
        if lat is not None and lng is not None:
            geo = {"lat": lat, "lng": lng, "barrio": None, "source": "excel", "from_cache": True}
        elif direccion:
            geo = geocode(direccion, cache, allow_net=not args.no_geocode)
        else:
            geo = {"lat": None, "lng": None, "barrio": None, "source": "sin-direccion"}
        rec["lat"], rec["lng"] = geo.get("lat"), geo.get("lng")
        rec["geocode_source"] = geo.get("source")
        rec["geocode_ok"] = rec["lat"] is not None
        rec["barrio"] = rec.get("barrio") or geo.get("barrio")
        rec["geo_display"] = geo.get("display")

        rec["maps_url"] = ("https://www.google.com/maps/search/?api=1&query="
                           + urllib.parse.quote(direccion)) if direccion else None

        # subte más cercana
        st, dist = nearest_subte(rec["lat"], rec["lng"], stations)
        rec["subte_nombre"] = st["name"] if st else None
        rec["subte_dist_m"] = round(dist) if dist is not None else None
        rec["subte_caminata_min"] = round(dist / 80.0) if dist is not None else None  # ~80 m/min

        # derivados (cuidando 0 legítimos -> chequeo is not None)
        precio, m2, pmax = rec.get("precio"), rec.get("m2"), rec.get("precio_max")
        neg = rec.get("negociar_a")
        if rec.get("usd_m2") is None and precio is not None and m2:
            rec["usd_m2"] = round(precio / m2, 1)
        if rec.get("usd_m2_max") is None and pmax is not None and m2:
            rec["usd_m2_max"] = round(pmax / m2, 1)
        if rec.get("p25") is None and precio is not None:
            rec["p25"] = round(precio * 0.25)
        if rec.get("p75") is None and precio is not None:
            rec["p75"] = round(precio * 0.75)
        # margen válido sólo si el target de negociación es un precio sensato (0 < neg < precio)
        rec["margen_neg_pct"] = (round((precio - neg) / precio * 100, 1)
                                 if precio and neg is not None and 0 < neg < precio else None)
        rec["dif_vs_tope"] = (round(precio - pmax)
                              if precio is not None and pmax is not None else None)
        rec["costo_mensual_usd"] = (round(rec["expensas"] / args.rate)
                                    if rec.get("expensas") is not None and args.rate else None)

        # flags de texto
        motivo = rec.get("motivo") or ""
        rec["refacciones"] = bool(RE_REFACCION.search(motivo))
        rec["transporte_motivo"] = bool(RE_TRANSPORTE.search(motivo))

        # flags de datos faltantes
        rec["faltantes"] = [lbl for f, lbl in
                            [("expensas", "Expensas"), ("antiguedad", "Antigüedad"),
                             ("orientacion", "Orientación"), ("posicion", "Posición")]
                            if rec.get(f) in (None, "")]
        props.append(rec)

    if not props:
        sys.exit("No encontré filas de datos debajo de los encabezados.")

    # ---- normalización para sub-scores (a nivel del conjunto) ----
    usd_m2_vals = [p.get("usd_m2") for p in props]
    precio_vals = [p.get("precio") for p in props]
    m2_vals = [p.get("m2") for p in props]
    amb_vals = [p.get("ambientes") for p in props]
    antig_vals = [p.get("antiguedad") for p in props]
    dist_vals = [p.get("subte_dist_m") for p in props]

    n_usd_m2 = minmax_norm(usd_m2_vals, invert=True)
    n_precio = minmax_norm(precio_vals, invert=True)
    n_m2 = minmax_norm(m2_vals)
    n_amb = minmax_norm(amb_vals)
    n_antig = minmax_norm(antig_vals, invert=True)
    n_dist = minmax_norm(dist_vals, invert=True)

    for i, p in enumerate(props):
        orient = ORIENT_SCORE.get(norm(p.get("orientacion")), 0.5)
        outdoor = OUTDOOR_SCORE.get(norm(p.get("balcon")), 0.4)
        posic = POSICION_SCORE.get(norm(p.get("posicion")), 0.6)

        s_precio = 0.6 * n_usd_m2[i] + 0.4 * n_precio[i]
        s_ubic = 0.8 * n_dist[i] + 0.2 * (1.0 if p["transporte_motivo"] else 0.0)
        s_tam = 0.7 * n_m2[i] + 0.3 * n_amb[i]
        s_estado = 0.35 * n_antig[i] + 0.25 * orient + 0.25 * outdoor + 0.15 * posic
        if p["refacciones"]:
            s_estado *= 0.85

        p["scores"] = {"precio": round(s_precio, 4), "ubicacion": round(s_ubic, 4),
                       "tamano": round(s_tam, 4), "estado": round(min(s_estado, 1.0), 4)}
        p["score_parts"] = {
            "n_usd_m2": round(n_usd_m2[i], 3), "n_precio": round(n_precio[i], 3),
            "n_dist": round(n_dist[i], 3), "n_m2": round(n_m2[i], 3),
            "n_amb": round(n_amb[i], 3), "n_antig": round(n_antig[i], 3),
            "orient": orient, "outdoor": outdoor, "posic": posic}

    # ---- estadísticas agregadas ----
    def stat(vals, fn):
        xs = [v for v in vals if v is not None]
        return fn(xs) if xs else None

    barrios = {}
    for p in props:
        b = p.get("barrio") or "Sin dato"
        barrios[b] = barrios.get(b, 0) + 1

    stats = {
        "n": len(props),
        "precio_min": stat(precio_vals, min), "precio_max": stat(precio_vals, max),
        "precio_prom": round(stat(precio_vals, lambda x: sum(x) / len(x))) if any(v is not None for v in precio_vals) else None,
        "usd_m2_prom": round(stat(usd_m2_vals, lambda x: sum(x) / len(x)), 1) if any(v is not None for v in usd_m2_vals) else None,
        "usd_m2_min": stat(usd_m2_vals, min), "usd_m2_max": stat(usd_m2_vals, max),
        "m2_prom": round(stat(m2_vals, lambda x: sum(x) / len(x)), 1) if any(v is not None for v in m2_vals) else None,
        "expensas_prom": round(stat([p.get("expensas") for p in props], lambda x: sum(x) / len(x)))
                         if any(p.get("expensas") is not None for p in props) else None,
        "antig_prom": round(stat(antig_vals, lambda x: sum(x) / len(x)), 1) if any(v is not None for v in antig_vals) else None,
        "apto_credito_n": sum(1 for p in props if norm(p.get("apto_credito")) in ("si", "yes", "true")),
        "barrios": barrios,
        "geocode_ok_n": sum(1 for p in props if p["geocode_ok"]),
    }

    payload = {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_file": excel.name,
        "usd_ars_rate": args.rate,
        "weights_default": {"precio": 0.25, "ubicacion": 0.25, "tamano": 0.25, "estado": 0.25},
        "fields_detected": sorted(mapping.keys()),
        "subte_stations": [{"name": s["name"], "lat": s["lat"], "lng": s["lon"]} for s in stations],
        "subte_lines": subte_lines,
        "stats": stats,
        "properties": props,
    }

    CACHE_FILE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    def emit(pl):
        return ("// AUTO-GENERADO por build.py — no editar a mano.\n"
                "window.__APP_DATA__ = " + json.dumps(pl, ensure_ascii=False, default=str) + ";\n")

    # data.js: SIEMPRE sin precios de negociación -> seguro para commitear y publicar.
    public = copy.deepcopy(payload)
    public["redacted_fields"] = SENSITIVE_FIELDS
    for p in public["properties"]:
        for k in SENSITIVE_FIELDS:
            if k in p:
                p[k] = None
    OUT_FILE.write_text(emit(public), encoding="utf-8")
    msg = f"OK -> {OUT_FILE.name} (público, sin precios de negociación)"

    # data.private.js: datos completos, sólo si se pide --full (gitignoreado, uso local).
    if args.full:
        PRIVATE_FILE.write_text(emit(payload), encoding="utf-8")
        msg += f"  +  {PRIVATE_FILE.name} (completo, local)"
    print(f"{msg}  ·  {len(props)} props, {stats['geocode_ok_n']} geolocalizadas")


if __name__ == "__main__":
    main()
